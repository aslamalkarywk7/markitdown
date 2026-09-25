"""Currency-formatted Excel cells keep their label (microsoft/markitdown#53)."""

import io

import openpyxl
import pytest

from markitdown import StreamInfo
from markitdown.converters import XlsxConverter

_INFO = StreamInfo(extension=".xlsx")


def _workbook() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["Item", "Count", "Cost", "Weight", "Total"])
    sheet.append(["Breakfast", 20, 5, None, 100])
    sheet.append(["Laptops", 5, 1199, None, 5995])
    sheet.append(["Car tires", 8, 199, "150 kg", 1592])
    for row in sheet.iter_rows(min_row=2, min_col=3, max_col=3):
        for cell in row:
            cell.number_format = '"$"#,##0.00'
    for row in sheet.iter_rows(min_row=2, min_col=5, max_col=5):
        for cell in row:
            cell.number_format = '"$"#,##0.00'
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def _convert(data: bytes) -> str:
    return XlsxConverter().convert(io.BytesIO(data), _INFO).markdown


def test_currency_cells_keep_their_label() -> None:
    markdown = _convert(_workbook())
    assert "$5" in markdown
    assert "$1199" in markdown
    assert "$100" in markdown
    assert "Breakfast" in markdown
    assert "150 kg" in markdown


def test_plain_cells_are_untouched() -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Item", "Count"])
    sheet.append(["Breakfast", 20])
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    markdown = _convert(stream.getvalue())
    assert "| Breakfast | 20 |" in markdown
    assert "$" not in markdown


def test_euro_suffix_format() -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Price"])
    sheet.append([42])
    sheet["A2"].number_format = "#,##0.00 [$€-x-euro2]"
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    assert "42€" in _convert(stream.getvalue())


def test_section_specific_currency_uses_cell_value() -> None:
    from markitdown.converters._xlsx_converter import (
        _currency_symbol,
        _is_currency_position_prefix,
        _select_format_section,
    )

    fmt = '"$"#,##0;"€"#,##0'
    assert _select_format_section(fmt, 5) == '"$"#,##0'
    assert _select_format_section(fmt, -5) == '"€"#,##0'
    assert _currency_symbol(fmt, 5) == "$"
    assert _currency_symbol(fmt, -5) == "€"
    assert _is_currency_position_prefix(fmt, 5)
    assert _is_currency_position_prefix(fmt, -5)

    # 3 sections: positive / negative / zero
    fmt3 = '"$"#,##0;"€"#,##0;"¥"#,##0'
    assert _currency_symbol(fmt3, 5) == "$"
    assert _currency_symbol(fmt3, -5) == "€"
    assert _currency_symbol(fmt3, 0) == "¥"

    # end-to-end: negative keeps its own section currency
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Balance"])
    sheet.append([10])
    sheet.append([-5])
    sheet["A2"].number_format = fmt
    sheet["A3"].number_format = fmt
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    markdown = _convert(stream.getvalue())
    assert "$10" in markdown
    assert "€" in markdown
    assert "$-5" not in markdown


def test_thai_baht_format() -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Price"])
    sheet.append([42])
    sheet["A2"].number_format = '"฿"#,##0'
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    assert "฿42" in _convert(stream.getvalue())


def test_quoted_semicolon_is_not_a_section_separator() -> None:
    from markitdown.converters._xlsx_converter import _select_format_section

    assert _select_format_section('"$;gross"#,##0', -5) == '"$;gross"#,##0'
